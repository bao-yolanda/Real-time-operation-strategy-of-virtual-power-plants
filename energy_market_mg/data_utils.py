from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import numpy as np
import openpyxl
import pandas as pd
import scipy.io as sio


@dataclass(frozen=True)
class EVType:
    name: str
    charge_power_mw: float
    energy_demand_mwh: float


EV_TYPES: tuple[EVType, ...] = (
    EVType("a", charge_power_mw=0.011, energy_demand_mwh=0.0375),
    EVType("b", charge_power_mw=0.022, energy_demand_mwh=0.0300),
    EVType("c", charge_power_mw=0.007, energy_demand_mwh=0.0225),
)


@dataclass
class EVTask:
    arrival: int
    departure: int
    charge_power_mw: float
    energy_demand_mwh: float


@dataclass
class MicrogridCase:
    horizon: int
    delta_t: float
    energy_price: np.ndarray
    sell_price: np.ndarray
    pv_forecast: np.ndarray
    pv_actual: np.ndarray
    load_forecast: np.ndarray
    load_actual: np.ndarray
    pv_floor_dayahead: np.ndarray
    load_ceiling_dayahead: np.ndarray
    ev_power_floor_dayahead: np.ndarray
    ev_cum_demand_ceiling_dayahead: np.ndarray
    ev_power_actual: np.ndarray
    ev_cum_demand_actual: np.ndarray
    grid_limit_mw: float
    es_power_ch_max_mw: float
    es_power_dis_max_mw: float
    es_energy_min_mwh: float
    es_energy_max_mwh: float
    es_energy_init_mwh: float
    eta_ch: float
    eta_dis: float
    deg_cost_mwh: float
    curtail_penalty_mwh: float
    ramp_penalty: float
    shortage_penalty_mwh: float
    metadata: dict


def _project_root() -> Path:
    return Path(__file__).resolve().parent.parent


def load_energy_price() -> np.ndarray:
    data = pd.read_csv(_project_root() / "market_price_data.csv")
    return data["能量价格"].to_numpy(dtype=float)


def load_pv_profile() -> np.ndarray:
    data = sio.loadmat(_project_root() / "data_prepare" / "output_pv.mat")
    return np.asarray(data["output_pv"]).flatten().astype(float)


def load_base_load(scale_mw: float = 1.8) -> np.ndarray:
    data = sio.loadmat(_project_root() / "data_prepare" / "h_load.mat")
    raw = np.asarray(data["h_load"], dtype=float)
    normalized = raw[:, 1]
    return normalized * scale_mw


def load_ev_tasks(
    ev_count: int = 40,
    seed: int = 42,
    arrival_uncertainty_h: int = 1,
    departure_uncertainty_h: int = 1,
    demand_uncertainty_ratio: float = 0.15,
) -> tuple[list[EVTask], list[EVTask], dict]:
    workbook = openpyxl.load_workbook(_project_root() / "data_prepare" / "EV_arrive_leave.xlsx", data_only=True)
    sheet = workbook["EV_arrive_leave"]

    rows: list[tuple[int, int]] = []
    row_idx = 2
    while True:
        arrive = sheet.cell(row=row_idx, column=2).value
        depart = sheet.cell(row=row_idx, column=3).value
        if arrive is None and depart is None:
            break
        if arrive is not None and depart is not None:
            rows.append((int(arrive) - 1, int(depart) - 1))
        row_idx += 1
    workbook.close()

    if not rows:
        raise ValueError("No EV arrival/departure data found.")

    rng = np.random.default_rng(seed)
    chosen = rows[: min(ev_count, len(rows))]
    forecast_tasks: list[EVTask] = []
    actual_tasks: list[EVTask] = []

    for idx, (arrival, departure) in enumerate(chosen):
        ev_type = EV_TYPES[idx % len(EV_TYPES)]
        base_arrival = int(np.clip(arrival, 0, 23))
        base_departure = int(np.clip(max(departure, base_arrival), base_arrival, 23))
        forecast_tasks.append(
            EVTask(
                arrival=base_arrival,
                departure=base_departure,
                charge_power_mw=ev_type.charge_power_mw,
                energy_demand_mwh=ev_type.energy_demand_mwh,
            )
        )

        arr_shift = int(rng.integers(-arrival_uncertainty_h, arrival_uncertainty_h + 1))
        dep_shift = int(rng.integers(-departure_uncertainty_h, departure_uncertainty_h + 1))
        demand_scale = float(rng.uniform(1.0 - demand_uncertainty_ratio, 1.0 + demand_uncertainty_ratio))
        actual_arrival = int(np.clip(base_arrival + arr_shift, 0, 23))
        actual_departure = int(np.clip(max(base_departure + dep_shift, actual_arrival), actual_arrival, 23))
        actual_tasks.append(
            EVTask(
                arrival=actual_arrival,
                departure=actual_departure,
                charge_power_mw=ev_type.charge_power_mw,
                energy_demand_mwh=ev_type.energy_demand_mwh * demand_scale,
            )
        )

    metadata = {
        "ev_count": len(chosen),
        "arrival_uncertainty_h": arrival_uncertainty_h,
        "departure_uncertainty_h": departure_uncertainty_h,
        "demand_uncertainty_ratio": demand_uncertainty_ratio,
    }
    return forecast_tasks, actual_tasks, metadata


def tasks_to_profiles(tasks: Iterable[EVTask], horizon: int) -> tuple[np.ndarray, np.ndarray]:
    available_power = np.zeros(horizon)
    cumulative_due = np.zeros(horizon)

    tasks_list = list(tasks)
    for t in range(horizon):
        available_power[t] = sum(task.charge_power_mw for task in tasks_list if task.arrival <= t <= task.departure)
        cumulative_due[t] = sum(task.energy_demand_mwh for task in tasks_list if task.departure <= t)

    cumulative_due = np.maximum.accumulate(cumulative_due)
    return available_power, cumulative_due


def sample_dro_envelope(
    pv_forecast: np.ndarray,
    load_forecast: np.ndarray,
    forecast_ev_tasks: list[EVTask],
    n_samples: int = 250,
    alpha: float = 0.15,
    pv_error_ratio: float = 0.18,
    load_error_ratio: float = 0.08,
    ev_time_shift_h: int = 1,
    ev_demand_ratio: float = 0.15,
    seed: int = 123,
) -> dict:
    horizon = len(pv_forecast)
    rng = np.random.default_rng(seed)

    pv_samples = np.zeros((n_samples, horizon))
    load_samples = np.zeros((n_samples, horizon))
    ev_power_samples = np.zeros((n_samples, horizon))
    ev_due_samples = np.zeros((n_samples, horizon))

    daytime_mask = (pv_forecast > 1e-6).astype(float)

    for s in range(n_samples):
        pv_noise = rng.normal(0.0, pv_error_ratio, size=horizon) * daytime_mask
        load_noise = rng.normal(0.0, load_error_ratio, size=horizon)
        pv_samples[s] = np.clip(pv_forecast * (1.0 + pv_noise), 0.0, None)
        load_samples[s] = np.clip(load_forecast * (1.0 + load_noise), 0.0, None)

        sampled_tasks: list[EVTask] = []
        for task in forecast_ev_tasks:
            arr_shift = int(rng.integers(-ev_time_shift_h, ev_time_shift_h + 1))
            dep_shift = int(rng.integers(-ev_time_shift_h, ev_time_shift_h + 1))
            demand_scale = float(rng.uniform(1.0 - ev_demand_ratio, 1.0 + ev_demand_ratio))
            arrival = int(np.clip(task.arrival + arr_shift, 0, horizon - 1))
            departure = int(np.clip(max(task.departure + dep_shift, arrival), arrival, horizon - 1))
            sampled_tasks.append(
                EVTask(
                    arrival=arrival,
                    departure=departure,
                    charge_power_mw=task.charge_power_mw,
                    energy_demand_mwh=task.energy_demand_mwh * demand_scale,
                )
            )
        ev_power_samples[s], ev_due_samples[s] = tasks_to_profiles(sampled_tasks, horizon)

    low_q = alpha
    high_q = 1.0 - alpha
    pv_floor = np.quantile(pv_samples, low_q, axis=0)
    load_ceiling = np.quantile(load_samples, high_q, axis=0)
    ev_power_floor = np.quantile(ev_power_samples, low_q, axis=0)
    ev_cum_demand_ceiling = np.quantile(ev_due_samples, high_q, axis=0)
    ev_cum_demand_ceiling = np.maximum.accumulate(ev_cum_demand_ceiling)

    return {
        "pv_floor": pv_floor,
        "load_ceiling": load_ceiling,
        "ev_power_floor": ev_power_floor,
        "ev_cum_demand_ceiling": ev_cum_demand_ceiling,
        "n_samples": n_samples,
        "alpha": alpha,
    }


def build_case(
    ev_count: int = 40,
    seed: int = 42,
    load_scale_mw: float = 1.8,
    grid_limit_mw: float = 3.5,
) -> MicrogridCase:
    horizon = 24
    delta_t = 1.0
    energy_price = load_energy_price()
    sell_price = 0.6 * energy_price
    pv_forecast = load_pv_profile()
    load_forecast = load_base_load(scale_mw=load_scale_mw)

    forecast_tasks, actual_tasks, ev_meta = load_ev_tasks(ev_count=ev_count, seed=seed)
    ev_power_forecast, ev_cum_demand_forecast = tasks_to_profiles(forecast_tasks, horizon)
    ev_power_actual, ev_cum_demand_actual = tasks_to_profiles(actual_tasks, horizon)

    rng = np.random.default_rng(seed + 7)
    pv_actual = np.clip(pv_forecast * (1.0 + rng.normal(0.0, 0.12, size=horizon) * (pv_forecast > 1e-6)), 0.0, None)
    load_actual = np.clip(load_forecast * (1.0 + rng.normal(0.0, 0.05, size=horizon)), 0.0, None)

    dro = sample_dro_envelope(
        pv_forecast=pv_forecast,
        load_forecast=load_forecast,
        forecast_ev_tasks=forecast_tasks,
        seed=seed + 99,
    )

    metadata = {
        "ev_forecast_total_demand_mwh": float(ev_cum_demand_forecast[-1]),
        "ev_actual_total_demand_mwh": float(ev_cum_demand_actual[-1]),
        "dro_samples": dro["n_samples"],
        "dro_alpha": dro["alpha"],
        **ev_meta,
    }

    return MicrogridCase(
        horizon=horizon,
        delta_t=delta_t,
        energy_price=energy_price,
        sell_price=sell_price,
        pv_forecast=pv_forecast,
        pv_actual=pv_actual,
        load_forecast=load_forecast,
        load_actual=load_actual,
        pv_floor_dayahead=dro["pv_floor"],
        load_ceiling_dayahead=dro["load_ceiling"],
        ev_power_floor_dayahead=dro["ev_power_floor"],
        ev_cum_demand_ceiling_dayahead=dro["ev_cum_demand_ceiling"],
        ev_power_actual=ev_power_actual,
        ev_cum_demand_actual=ev_cum_demand_actual,
        grid_limit_mw=grid_limit_mw,
        es_power_ch_max_mw=1.0,
        es_power_dis_max_mw=1.0,
        es_energy_min_mwh=0.2,
        es_energy_max_mwh=1.8,
        es_energy_init_mwh=1.0,
        eta_ch=0.95,
        eta_dis=0.95,
        deg_cost_mwh=12.0,
        curtail_penalty_mwh=40.0,
        ramp_penalty=4.0,
        shortage_penalty_mwh=3000.0,
        metadata=metadata,
    )

